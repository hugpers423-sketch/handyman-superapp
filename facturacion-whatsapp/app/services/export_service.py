import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill, numbers
from openpyxl.utils import get_column_letter
import io
import logging
import csv
from typing import List, Dict, Any, Optional
from sqlalchemy import select, func, and_
from sqlalchemy.ext.asyncio import AsyncSession
from datetime import datetime, date, timedelta

from app.models import User, Invoice, InvoiceItem, Product, Customer, Expense

logger = logging.getLogger(__name__)


class ExportService:
    """Exportar datos a Excel, CSV, Google Sheets"""
    
    def __init__(self, db: AsyncSession):
        self.db = db

    # ===== EXCEL EXPORTS =====
    
    async def export_invoices_excel(
        self,
        user_id: int,
        start_date: date = None,
        end_date: date = None,
        status: str = None
    ) -> bytes:
        """Exportar facturas a Excel"""
        wb = openpyxl.Workbook()
        
        # Hoja 1: Resumen
        ws1 = wb.active
        ws1.title = "Resumen"
        self._write_summary_sheet(ws1, user_id, start_date, end_date)
        
        # Hoja 2: Facturas detalle
        ws2 = wb.create_sheet("Facturas")
        await self._write_invoices_sheet(ws2, user_id, start_date, end_date, status)
        
        # Hoja 3: Items
        ws3 = wb.create_sheet("Items")
        await self._write_items_sheet(ws3, user_id, start_date, end_date)
        
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output.getvalue()

    async def export_products_excel(self, user_id: int) -> bytes:
        """Exportar catálogo de productos a Excel"""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Productos"
        
        header_font = Font(bold=True, color="FFFFFF", size=11)
        header_fill = PatternFill(start_color="0284C7", end_color="0284C7", fill_type="solid")
        thin_border = Border(
            left=Side(style="thin"), right=Side(style="thin"),
            top=Side(style="thin"), bottom=Side(style="thin")
        )
        
        headers = [
            "ID", "Nombre", "Descripción", "SKU", "Código Barras",
            "Categoría", "Marca", "Precio", "Stock", "Unidad",
            "Tipo IGV", "Tasa IGV", "Activo", "Creado"
        ]
        
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")
            cell.border = thin_border
        
        result = await self.db.execute(
            select(Product)
            .where(Product.user_id == user_id, Product.is_active == True)
            .order_by(Product.name)
        )
        products = result.scalars().all()
        
        money_format = '#,##0.00'
        for row_num, prod in enumerate(products, 2):
            row_data = [
                prod.id, prod.name, prod.description or "", prod.sku or "",
                prod.barcode or "", prod.category or "", prod.brand or "",
                prod.price, prod.stock, prod.unit,
                prod.tax_type, prod.tax_rate, "Sí" if prod.is_active else "No",
                prod.created_at.strftime("%d/%m/%Y %H:%M") if prod.created_at else ""
            ]
            
            for col, value in enumerate(row_data, 1):
                cell = ws.cell(row=row_num, column=col, value=value)
                cell.border = thin_border
                if col == 8:  # Precio
                    cell.number_format = money_format
        
        for col in range(1, len(headers) + 1):
            ws.column_dimensions[get_column_letter(col)].width = 18
        
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output.getvalue()

    async def export_customers_excel(self, user_id: int) -> bytes:
        """Exportar clientes a Excel"""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Clientes"
        
        header_font = Font(bold=True, color="FFFFFF", size=11)
        header_fill = PatternFill(start_color="16A34A", end_color="16A34A", fill_type="solid")
        thin_border = Border(
            left=Side(style="thin"), right=Side(style="thin"),
            top=Side(style="thin"), bottom=Side(style="thin")
        )
        
        headers = [
            "ID", "Nombre", "Nombre Comercial", "Tipo Doc", "Número Doc",
            "Teléfono", "Email", "Dirección", "Distrito", "Provincia", "Departamento",
            "WhatsApp", "Límite Crédito", "Crédito Usado", "Crédito Disponible",
            "Días Crédito", "Total Compras", "N° Facturas", "Última Compra",
            "Primera Compra", "Bloqueado", "Activo", "Creado"
        ]
        
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", wrap_text=True)
            cell.border = thin_border
        
        result = await self.db.execute(
            select(Customer).where(Customer.user_id == user_id).order_by(Customer.name)
        )
        customers = result.scalars().all()
        
        money_format = '#,##0.00'
        for row_num, cust in enumerate(customers, 2):
            available = cust.credit_limit - cust.credit_used if cust.credit_limit > 0 else 0
            row_data = [
                cust.id, cust.name, cust.trade_name or "", cust.doc_type, cust.doc_number,
                cust.phone or "", cust.email or "", cust.address or "",
                cust.district or "", cust.province or "", cust.department or "",
                cust.whatsapp_id or "", cust.credit_limit, cust.credit_used, available,
                cust.credit_days, cust.total_purchases, cust.total_invoices,
                cust.last_purchase_date.strftime("%d/%m/%Y") if cust.last_purchase_date else "",
                cust.first_purchase_date.strftime("%d/%m/%Y") if cust.first_purchase_date else "",
                "Sí" if cust.is_blocked else "No", "Sí" if cust.is_active else "No",
                cust.created_at.strftime("%d/%m/%Y %H:%M") if cust.created_at else ""
            ]
            
            for col, value in enumerate(row_data, 1):
                cell = ws.cell(row=row_num, column=col, value=value)
                cell.border = thin_border
                if col in [13, 14, 15, 17]:
                    cell.number_format = money_format
        
        for col in range(1, len(headers) + 1):
            ws.column_dimensions[get_column_letter(col)].width = 20
        
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output.getvalue()

    async def export_expenses_excel(
        self,
        user_id: int,
        year: int = None,
        month: int = None
    ) -> bytes:
        """Exportar gastos a Excel"""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Gastos"
        
        header_font = Font(bold=True, color="FFFFFF", size=11)
        header_fill = PatternFill(start_color="EF4444", end_color="EF4444", fill_type="solid")
        thin_border = Border(
            left=Side(style="thin"), right=Side(style="thin"),
            top=Side(style="thin"), bottom=Side(style="thin")
        )
        
        headers = [
            "Fecha", "Categoría", "Proveedor", "RUC Proveedor",
            "Tipo Doc", "Serie", "Número", "Descripción",
            "Subtotal", "IGV", "Total", "Moneda",
            "Método Pago", "Estado Pago", "Estado SUNAT"
        ]
        
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", wrap_text=True)
            cell.border = thin_border
        
        query = select(Expense).where(Expense.user_id == user_id)
        if year and month:
            start = date(year, month, 1)
            if month == 12:
                end = date(year + 1, 1, 1) - timedelta(days=1)
            else:
                end = date(year, month + 1, 1) - timedelta(days=1)
            query = query.where(
                func.date(Expense.expense_date) >= start,
                func.date(Expense.expense_date) <= end
            )
        query = query.order_by(Expense.expense_date.desc())
        
        result = await self.db.execute(query)
        expenses = result.scalars().all()
        
        money_format = '#,##0.00'
        for row_num, exp in enumerate(expenses, 2):
            row_data = [
                exp.expense_date.strftime("%d/%m/%Y"),
                exp.category or "", exp.supplier_name or "", exp.supplier_doc or "",
                exp.document_type or "", exp.series or "", exp.number or 0,
                exp.description or "",
                exp.subtotal, exp.tax_amount, exp.total, exp.currency,
                exp.payment_method, exp.payment_status, exp.sunat_status
            ]
            
            for col, value in enumerate(row_data, 1):
                cell = ws.cell(row=row_num, column=col, value=value)
                cell.border = thin_border
                if col in [9, 10, 11]:
                    cell.number_format = money_format
        
        for col in range(1, len(headers) + 1):
            ws.column_dimensions[get_column_letter(col)].width = 18
        
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output.getvalue()

    async def export_full_backup_excel(self, user_id: int) -> bytes:
        """Backup completo: todas las hojas en un archivo"""
        wb = openpyxl.Workbook()
        
        # Eliminar hoja por defecto
        wb.remove(wb.active)
        
        # Crear todas las hojas
        ws1 = wb.create_sheet("Facturas")
        await self._write_invoices_sheet(ws1, user_id)
        
        ws2 = wb.create_sheet("Items Facturas")
        await self._write_items_sheet(ws2, user_id)
        
        ws3 = wb.create_sheet("Productos")
        await self._write_products_sheet(ws3, user_id)
        
        ws4 = wb.create_sheet("Clientes")
        await self._write_customers_sheet(ws4, user_id)
        
        ws5 = wb.create_sheet("Gastos")
        await self._write_expenses_sheet(ws5, user_id)
        
        ws6 = wb.create_sheet("Resumen Mensual")
        await self._write_monthly_summary_sheet(ws6, user_id)
        
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output.getvalue()

    # ===== HELPER METHODS =====
    
    def _write_summary_sheet(self, ws, user_id: int, start_date: date = None, end_date: date = None):
        header_font = Font(bold=True, color="FFFFFF", size=11)
        header_fill = PatternFill(start_color="7C3AED", end_color="7C3AED", fill_type="solid")
        thin_border = Border(
            left=Side(style="thin"), right=Side(style="thin"),
            top=Side(style="thin"), bottom=Side(style="thin")
        )
        
        ws.cell(row=1, column=1, value="REPORTE DE FACTURACIÓN").font = Font(bold=True, size=14)
        ws.cell(row=2, column=1, value=f"Generado: {datetime.now().strftime('%d/%m/%Y %H:%M')}")
        ws.cell(row=3, column=1, value=f"Período: {start_date or 'Inicio'} al {end_date or 'Hoy'}")
        
        # TODO: Agregar métricas calculadas

    async def _write_invoices_sheet(self, ws, user_id: int, start_date: date = None, end_date: date = None, status: str = None):
        header_font = Font(bold=True, color="FFFFFF", size=10)
        header_fill = PatternFill(start_color="0284C7", end_color="0284C7", fill_type="solid")
        thin_border = Border(
            left=Side(style="thin"), right=Side(style="thin"),
            top=Side(style="thin"), bottom=Side(style="thin")
        )
        
        headers = [
            "Fecha", "Tipo", "Serie", "Número", "Cliente", "Doc Cliente",
            "Subtotal", "IGV", "Total", "Moneda", "Estado SUNAT", "PDF", "XML"
        ]
        
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")
            cell.border = thin_border
        
        query = select(Invoice).where(Invoice.user_id == user_id)
        if start_date:
            query = query.where(func.date(Invoice.created_at) >= start_date)
        if end_date:
            query = query.where(func.date(Invoice.created_at) <= end_date)
        if status:
            query = query.where(Invoice.sunat_status == status)
        query = query.order_by(Invoice.created_at.desc())
        
        result = await self.db.execute(query)
        invoices = result.scalars().all()
        
        money_format = '#,##0.00'
        for row_num, inv in enumerate(invoices, 2):
            row_data = [
                inv.created_at.strftime("%d/%m/%Y %H:%M"),
                inv.document_type_name, inv.series, inv.number,
                inv.customer_name or "", inv.customer_doc or "",
                inv.subtotal, inv.tax_amount, inv.total, inv.currency,
                inv.sunat_status, inv.pdf_url or "", inv.xml_url or ""
            ]
            
            for col, value in enumerate(row_data, 1):
                cell = ws.cell(row=row_num, column=col, value=value)
                cell.border = thin_border
                if col in [7, 8, 9]:
                    cell.number_format = money_format
        
        for col in range(1, len(headers) + 1):
            ws.column_dimensions[get_column_letter(col)].width = 20

    async def _write_items_sheet(self, ws, user_id: int, start_date: date = None, end_date: date = None):
        header_font = Font(bold=True, color="FFFFFF", size=10)
        header_fill = PatternFill(start_color="22C55E", end_color="22C55E", fill_type="solid")
        thin_border = Border(
            left=Side(style="thin"), right=Side(style="thin"),
            top=Side(style="thin"), bottom=Side(style="thin")
        )
        
        headers = [
            "Factura Fecha", "Factura Tipo", "Factura Serie", "Factura Número",
            "Item #", "Producto", "Cantidad", "Unidad", "P. Unitario",
            "Descuento %", "IGV %", "Total Item"
        ]
        
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")
            cell.border = thin_border
        
        query = (
            select(InvoiceItem, Invoice)
            .join(Invoice, InvoiceItem.invoice_id == Invoice.id)
            .where(Invoice.user_id == user_id)
        )
        if start_date:
            query = query.where(func.date(Invoice.created_at) >= start_date)
        if end_date:
            query = query.where(func.date(Invoice.created_at) <= end_date)
        query = query.order_by(Invoice.created_at.desc(), InvoiceItem.id)
        
        result = await self.db.execute(query)
        rows = result.all()
        
        money_format = '#,##0.00'
        for row_num, (item, inv) in enumerate(rows, 2):
            row_data = [
                inv.created_at.strftime("%d/%m/%Y"),
                inv.document_type_name, inv.series, inv.number,
                row_num - 1, item.description, item.quantity, item.unit,
                item.unit_price, item.discount, item.tax_rate, item.total
            ]
            
            for col, value in enumerate(row_data, 1):
                cell = ws.cell(row=row_num, column=col, value=value)
                cell.border = thin_border
                if col in [9, 12]:
                    cell.number_format = money_format
        
        for col in range(1, len(headers) + 1):
            ws.column_dimensions[get_column_letter(col)].width = 18

    async def _write_products_sheet(self, ws, user_id: int):
        # Similar a export_products_excel pero sin crear nuevo workbook
        await self._write_sheet_from_model(ws, Product, user_id, [
            "id", "name", "description", "sku", "barcode", "category", "brand",
            "price", "stock", "unit", "tax_type", "tax_rate", "is_active", "created_at"
        ])

    async def _write_customers_sheet(self, ws, user_id: int):
        await self._write_sheet_from_model(ws, Customer, user_id, [
            "id", "name", "trade_name", "doc_type", "doc_number", "phone", "email",
            "address", "district", "province", "department", "whatsapp_id",
            "credit_limit", "credit_used", "credit_days", "total_purchases",
            "total_invoices", "last_purchase_date", "first_purchase_date",
            "is_blocked", "is_active", "created_at"
        ])

    async def _write_expenses_sheet(self, ws, user_id: int):
        await self._write_sheet_from_model(ws, Expense, user_id, [
            "expense_date", "category", "supplier_name", "supplier_doc",
            "document_type", "series", "number", "description",
            "subtotal", "tax_amount", "total", "currency",
            "payment_method", "payment_status", "sunat_status"
        ])

    async def _write_sheet_from_model(self, ws, model, user_id: int, fields: List[str]):
        header_font = Font(bold=True, color="FFFFFF", size=10)
        header_fill = PatternFill(start_color="6B7280", end_color="6B7280", fill_type="solid")
        thin_border = Border(
            left=Side(style="thin"), right=Side(style="thin"),
            top=Side(style="thin"), bottom=Side(style="thin")
        )
        
        for col, field in enumerate(fields, 1):
            cell = ws.cell(row=1, column=col, value=field.replace("_", " ").title())
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")
            cell.border = thin_border
        
        result = await self.db.execute(
            select(model).where(model.user_id == user_id).order_by(model.id)
        )
        records = result.scalars().all()
        
        for row_num, record in enumerate(records, 2):
            for col, field in enumerate(fields, 1):
                value = getattr(record, field, "")
                if isinstance(value, datetime):
                    value = value.strftime("%d/%m/%Y %H:%M")
                elif isinstance(value, date):
                    value = value.strftime("%d/%m/%Y")
                elif isinstance(value, bool):
                    value = "Sí" if value else "No"
                
                cell = ws.cell(row=row_num, column=col, value=value)
                cell.border = thin_border
        
        for col in range(1, len(fields) + 1):
            ws.column_dimensions[get_column_letter(col)].width = 18

    async def _write_monthly_summary_sheet(self, ws, user_id: int):
        """Hoja con resumen por mes del año actual"""
        header_font = Font(bold=True, color="FFFFFF", size=10)
        header_fill = PatternFill(start_color="F59E0B", end_color="F59E0B", fill_type="solid")
        thin_border = Border(
            left=Side(style="thin"), right=Side(style="thin"),
            top=Side(style="thin"), bottom=Side(style="thin")
        )
        
        headers = ["Mes", "Ventas", "Facturas", "Gastos", "Utilidad", "Margen %", "IGV Neto"]
        
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")
            cell.border = thin_border
        
        current_year = datetime.now().year
        money_format = '#,##0.00'
        
        for month in range(1, 13):
            # Ventas
            sales_result = await self.db.execute(
                select(
                    func.sum(Invoice.total), func.count(Invoice.id),
                    func.sum(Invoice.subtotal), func.sum(Invoice.tax_amount)
                ).where(
                    Invoice.user_id == user_id,
                    extract('year', Invoice.created_at) == current_year,
                    extract('month', Invoice.created_at) == month,
                    Invoice.sunat_status == "accepted"
                )
            )
            sales_total, sales_count, sales_subtotal, sales_igv = sales_result.first()
            
            # Gastos
            exp_result = await self.db.execute(
                select(
                    func.sum(Expense.total), func.count(Expense.id),
                    func.sum(Expense.subtotal), func.sum(Expense.tax_amount)
                ).where(
                    Expense.user_id == user_id,
                    extract('year', Expense.expense_date) == current_year,
                    extract('month', Expense.expense_date) == month
                )
            )
            exp_total, exp_count, exp_subtotal, exp_igv = exp_result.first()
            
            sales_total = float(sales_total or 0)
            exp_total = float(exp_total or 0)
            profit = sales_total - exp_total
            margin = (profit / sales_total * 100) if sales_total > 0 else 0
            igv_net = float(sales_igv or 0) - float(exp_igv or 0)
            
            row_data = [
                f"{month:02d}/{current_year}",
                sales_total, sales_count or 0,
                exp_total, profit, margin, max(0, igv_net)
            ]
            
            row_num = month + 1
            for col, value in enumerate(row_data, 1):
                cell = ws.cell(row=row_num, column=col, value=value)
                cell.border = thin_border
                if col in [2, 4, 5, 7]:
                    cell.number_format = money_format
                elif col == 6:
                    cell.number_format = '0.00%'
                    cell.value = margin / 100
        
        for col in range(1, len(headers) + 1):
            ws.column_dimensions[get_column_letter(col)].width = 18

    # ===== CSV EXPORTS =====
    
    async def export_invoices_csv(
        self,
        user_id: int,
        start_date: date = None,
        end_date: date = None
    ) -> str:
        """Exportar facturas a CSV string"""
        output = io.StringIO()
        writer = csv.writer(output)
        
        writer.writerow([
            "Fecha", "Tipo", "Serie", "Numero", "Cliente", "DocCliente",
            "Subtotal", "IGV", "Total", "Moneda", "EstadoSUNAT"
        ])
        
        query = select(Invoice).where(Invoice.user_id == user_id)
        if start_date:
            query = query.where(func.date(Invoice.created_at) >= start_date)
        if end_date:
            query = query.where(func.date(Invoice.created_at) <= end_date)
        query = query.order_by(Invoice.created_at.desc())
        
        result = await self.db.execute(query)
        invoices = result.scalars().all()
        
        for inv in invoices:
            writer.writerow([
                inv.created_at.strftime("%Y-%m-%d"),
                inv.document_type, inv.series, inv.number,
                inv.customer_name or "", inv.customer_doc or "",
                inv.subtotal, inv.tax_amount, inv.total,
                inv.currency, inv.sunat_status
            ])
        
        return output.getvalue()

    # ===== GOOGLE SHEETS INTEGRATION (stub) =====
    
    async def sync_to_google_sheets(
        self,
        user_id: int,
        spreadsheet_id: str,
        credentials_json: Dict[str, Any],
        sheets_config: Dict[str, str]
    ) -> Dict[str, Any]:
        """
        Sincronizar datos a Google Sheets.
        Requiere: google-auth, google-api-python-client, gspread
        """
        # TODO: Implementar con gspread o google-api-python-client
        # Ejemplo de estructura:
        """
        import gspread
        from google.oauth2.service_account import Credentials
        
        scopes = ['https://www.googleapis.com/auth/spreadsheets']
        creds = Credentials.from_service_account_info(credentials_json, scopes=scopes)
        gc = gspread.authorize(creds)
        sh = gc.open_by_key(spreadsheet_id)
        
        # Sync facturas
        worksheet = sh.worksheet(sheets_config.get('invoices', 'Facturas'))
        data = await self.export_invoices_csv(user_id)
        # ... procesar y subir
        """
        
        return {
            "success": False,
            "error": "Integración Google Sheets pendiente. Instalar gspread y configurar credenciales."
        }


# ===== FUNCIONES DE CONVENIENCIA =====

async def generate_excel_backup(user_id: int) -> bytes:
    """Generar backup completo en Excel"""
    from app.core.database import AsyncSessionLocal
    
    async with AsyncSessionLocal() as db:
        service = ExportService(db)
        return await service.export_full_backup_excel(user_id)


async def send_excel_backup_whatsapp(
    whatsapp_id: str,
    user_id: int,
    whatsapp_service,
    filename: str = None
):
    """Enviar backup Excel por WhatsApp"""
    try:
        excel_bytes = await generate_excel_backup(user_id)
        
        # En producción: subir a S3/temporal y enviar URL
        # Por ahora solo notificar
        await whatsapp_service.send_text_message(
            whatsapp_id,
            f"📊 *Backup Excel generado*\n\n"
            f"El archivo está listo para descargar desde el dashboard.\n"
            f"Ve a /dashboard/settings para descargar."
        )
    except Exception as e:
        logger.error(f"Error generando backup: {e}")
        await whatsapp_service.send_text_message(
            whatsapp_id,
            f"❌ Error generando backup: {str(e)}"
        )